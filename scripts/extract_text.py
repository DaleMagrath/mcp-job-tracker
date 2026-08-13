"""
Extract plain text from a resume/cover-letter file for read_document.

Usage:  python extract_text.py <file.pdf|file.docx>
Prints a JSON object to stdout: {"text": ..., "words": N, "chars": N, "pages": N?}
Exits non-zero with a stderr message on unsupported type or a missing library
(the caller turns that into a friendly error).
"""
import json
import os
import sys


def extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = len(reader.pages)
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    return text, {"pages": pages}


def extract_docx(path):
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts), {}


def main():
    if len(sys.argv) < 2:
        sys.stderr.write("usage: extract_text.py <file>\n")
        return 2
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            text, meta = extract_pdf(path)
        elif ext == ".docx":
            text, meta = extract_docx(path)
        else:
            sys.stderr.write(f"unsupported file type: {ext or '(none)'}\n")
            return 2
    except ImportError as e:
        sys.stderr.write(f"missing extraction library: {e}\n")
        return 3
    except Exception as e:  # noqa: BLE001
        sys.stderr.write(f"could not extract text: {e}\n")
        return 4

    # Collapse runs of blank lines so the output stays readable.
    lines = [ln.rstrip() for ln in text.splitlines()]
    cleaned, blank = [], False
    for ln in lines:
        if ln.strip() == "":
            if not blank:
                cleaned.append("")
            blank = True
        else:
            cleaned.append(ln)
            blank = False
    text = "\n".join(cleaned).strip()

    out = {"text": text, "words": len(text.split()), "chars": len(text)}
    out.update(meta)
    sys.stdout.write(json.dumps(out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
