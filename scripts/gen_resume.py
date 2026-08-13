"""
Render a resume .docx from a JSON spec, matching the format Claude's prior
resume scripts used (2-page, Liberation Serif). Conversion to PDF is done by the
caller (LibreOffice).

Usage:  python gen_resume.py <input.json> <output.docx>

Key correctness note: role/date heading lines are rendered as a 2-column,
borderless TABLE (title left, dates right) — NOT a tab with a right tab stop
across mixed bold/italic runs. The tab construct triggered a LibreOffice PDF
ToUnicode-CMap bug on subset fonts that made the bold title extract as garbled
characters (pdftotext / pypdf / pdfminer). The table extracts cleanly. Keep it.

Input JSON shape (merged by the MCP tool from resume_master.json + tailoring):
{
  "name": "Dale Magrath",
  "contact": "email • location • linkedin.com/in/...",
  "summary": "tailored summary paragraph",
  "key_qualifications": { "target": "Acme Corp Engineering Manager", "bullets": [...] },
  "experience": [ { "org": "...", "title": "...", "dates": "...", "description": "..." } ],
  "projects": [ { "name": "...", "dates": "...", "description": "..." } ],   # optional
  "education": ["..."],
  "certifications": ["..."],
  "skills": ["...", "..."],
  "honors": ["..."]                                                          # optional, off by default
}
"""
import json
import sys

from docx import Document
from docx.shared import Pt, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HEADING_COLOR = RGBColor(0x2E, 0x40, 0x57)
CONTACT_COLOR = RGBColor(0x44, 0x44, 0x44)
DATES_COLOR = RGBColor(0x55, 0x55, 0x55)
LINK_COLOR = "1155CC"
FONT = "Liberation Serif"

# Content width = page (12240) - left (900) - right (900) = 10440 twips.
CONTENT_TWIPS = 10440
LEFT_COL_TWIPS = int(CONTENT_TWIPS * 0.70)
RIGHT_COL_TWIPS = CONTENT_TWIPS - LEFT_COL_TWIPS


def _bottom_border(p, sz, color, space):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)
    pbdr.append(b)
    pPr.append(pbdr)


def _add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), FONT)
    rf.set(qn("w:hAnsi"), FONT)
    rPr.append(rf)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), LINK_COLOR)
    rPr.append(col)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")  # 10pt
    rPr.append(sz)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _zero_cell_margins(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge in ("top", "start", "bottom", "end", "left", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def _no_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        borders.append(e)
    tblPr.append(borders)


def build(spec, out_path):
    doc = Document()

    for s in doc.sections:
        s.page_width = Twips(12240)
        s.page_height = Twips(15840)
        s.top_margin = Twips(720)
        s.bottom_margin = Twips(720)
        s.left_margin = Twips(900)
        s.right_margin = Twips(900)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(5)
        _bottom_border(p, 4, "CCCCCC", 2)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = HEADING_COLOR
        return p

    def role(title, dates):
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.allow_autofit = False
        _no_table_borders(table)
        table.columns[0].width = Twips(LEFT_COL_TWIPS)
        table.columns[1].width = Twips(RIGHT_COL_TWIPS)
        left, right = table.cell(0, 0), table.cell(0, 1)
        left.width = Twips(LEFT_COL_TWIPS)
        right.width = Twips(RIGHT_COL_TWIPS)
        _zero_cell_margins(left)
        _zero_cell_margins(right)

        lp = left.paragraphs[0]
        lp.paragraph_format.space_before = Pt(8)
        lp.paragraph_format.space_after = Pt(2)
        lr = lp.add_run(title)
        lr.bold = True
        lr.font.size = Pt(10.5)
        lr.font.name = FONT

        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.paragraph_format.space_before = Pt(8)
        rp.paragraph_format.space_after = Pt(2)
        rr = rp.add_run(dates)
        rr.italic = True
        rr.font.size = Pt(10)
        rr.font.color.rgb = DATES_COLOR
        rr.font.name = FONT

    def body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(text)
        r.font.size = Pt(10.5)
        return p

    def bullet(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Twips(230)
        p.paragraph_format.first_line_indent = Twips(-230)
        r = p.add_run("• " + text)
        r.font.size = Pt(10)
        return p

    # ---- Header ----
    p = doc.add_paragraph()
    r = p.add_run(spec.get("name", ""))
    r.bold = True
    r.font.size = Pt(20)

    contact = spec.get("contact", "")
    cp = doc.add_paragraph()
    cp.paragraph_format.space_after = Pt(6)
    _bottom_border(cp, 6, "999999", 6)
    parts = [t.strip() for t in contact.split("•")] if contact else []
    for i, tok in enumerate(parts):
        if i:
            sep = cp.add_run("  •  ")
            sep.font.size = Pt(10)
            sep.font.color.rgb = CONTACT_COLOR
        if "linkedin.com" in tok.lower():
            url = tok if tok.lower().startswith("http") else "https://" + tok
            _add_hyperlink(cp, url, tok)
        else:
            run = cp.add_run(tok)
            run.font.size = Pt(10)
            run.font.color.rgb = CONTACT_COLOR

    # ---- Summary ----
    if spec.get("summary"):
        heading("Summary")
        body(spec["summary"])

    # ---- Key qualifications (tailored alignment section) ----
    kq = spec.get("key_qualifications") or {}
    bullets = kq.get("bullets") or []
    if bullets:
        target = (kq.get("target") or "").strip()
        label = f"Key Qualifications — {target} Alignment" if target else "Key Qualifications"
        heading(label)
        for b in bullets:
            bullet(b)

    # ---- Experience ----
    for section, items, is_exp in (
        ("Experience", spec.get("experience") or [], True),
        ("Projects", spec.get("projects") or [], False),
    ):
        if not items:
            continue
        heading(section)
        for e in items:
            if is_exp:
                org, title = e.get("org", "").strip(), e.get("title", "").strip()
                head = f"{org}, {title}" if org and title else (org or title)
            else:
                head = e.get("name", "")
            role(head, e.get("dates", ""))
            if e.get("description"):
                body(e["description"])

    # ---- Education ----
    if spec.get("education"):
        heading("Education")
        for line in spec["education"]:
            body(line)

    # ---- Certifications ----
    if spec.get("certifications"):
        heading("Licenses & Certifications")
        for line in spec["certifications"]:
            body(line)

    # ---- Skills ----
    if spec.get("skills"):
        heading("Skills")
        body(" • ".join(spec["skills"]))

    # ---- Honors & Awards (optional, off by default for 2-page fit) ----
    if spec.get("honors"):
        heading("Honors & Awards")
        for line in spec["honors"]:
            body(line)

    doc.save(out_path)


def main():
    if len(sys.argv) < 3:
        sys.stderr.write("usage: python gen_resume.py <input.json> <output.docx>\n")
        return 2
    try:
        with open(sys.argv[1], "r", encoding="utf-8") as f:
            spec = json.load(f)
    except Exception as e:  # noqa: BLE001
        sys.stderr.write(f"could not read spec JSON: {e}\n")
        return 3
    try:
        build(spec, sys.argv[2])
    except ImportError:
        sys.stderr.write("python-docx not installed (pip install python-docx)\n")
        return 2
    print("WROTE", sys.argv[2])
    return 0


if __name__ == "__main__":
    sys.exit(main())
