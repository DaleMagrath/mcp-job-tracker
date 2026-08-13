"""
Re-apply the standard formatting to Job_Search_Discovery.xlsx.

The MCP server writes the discovery sheet with SheetJS (great for row data, but
the community build does not re-serialize rich cell styling). This script is run
after those writes so the file keeps the same look the daily openpyxl task gives
it: bold header, frozen top row, auto-filter, wrapped text, sensible column
widths, and clickable Job Link cells.

Usage:  python format_discovery.py "<path to .xlsx>"
Best-effort: exits non-zero with a message if openpyxl is missing.
"""

import sys


def main(path: str) -> int:
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.stderr.write("openpyxl not installed (pip install openpyxl)\n")
        return 2

    wb = load_workbook(path)
    ws = wb.worksheets[0]
    if ws.max_row < 1 or ws.max_column < 1:
        wb.save(path)
        return 0

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    # Header row: bold white on blue, wrapped, centered.
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Locate the "Job Link" column, if present.
    link_col = None
    for idx, h in enumerate(headers, start=1):
        if isinstance(h, str) and h.strip().lower() == "job link":
            link_col = idx
            break

    link_font = Font(color="0563C1", underline="single")
    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).alignment = Alignment(
                vertical="top", wrap_text=True
            )
        if link_col:
            lc = ws.cell(row=row, column=link_col)
            val = lc.value
            if isinstance(val, str) and val.strip().startswith(("http://", "https://")):
                lc.hyperlink = val.strip()
                lc.font = link_font

    # Column widths from content, clamped to a readable range.
    for col in range(1, ws.max_column + 1):
        longest = 0
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row=row, column=col).value
            if v is not None:
                longest = max(longest, min(len(str(v)), 60))
        ws.column_dimensions[get_column_letter(col)].width = max(
            12, min(longest + 2, 45)
        )

    wb.save(path)
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.stderr.write("usage: python format_discovery.py <path.xlsx>\n")
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
